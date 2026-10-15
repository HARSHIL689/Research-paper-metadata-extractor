import os
import re
from typing import Dict, Any, List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

EXCEL_FILE_PATH = os.path.join("data", "papers_metadata.xlsx")

COLUMNS = [
    "Source File", "Title", "Authors", "Publication Date", "Journal / Venue",
    "Volume", "DOI", "arXiv ID", "Keywords", "Language", "Quality Score", "Abstract",
]


def format_authors(authors: List[Dict[str, Any]]) -> str:
    """Return author names only. Emails/affiliations are intentionally excluded."""
    rows = []
    seen = set()
    for author in authors or []:
        name = " ".join(str(author.get("name") or "").split()).strip()
        if not name:
            continue
        # Defensive cleanup in case a parser accidentally puts contact data in the name.
        name = re.sub(r"\s*[<(]?[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}[)>]?", "", name, flags=re.I).strip()
        name = re.sub(r"\s*\([^)]*\b(?:Department|University|Institute|College|School|Laboratory|Lab|Faculty)\b[^)]*\)$", "", name, flags=re.I).strip()
        key = " ".join(name.casefold().split())
        if name and key not in seen:
            seen.add(key)
            rows.append(name)
    return "\n".join(rows)


def join_unique(values) -> str:
    result = []
    for value in values or []:
        if value and value not in result:
            result.append(str(value))
    return " ; ".join(result)


def flatten_metadata(metadata: Dict[str, Any]) -> Dict[str, Any]:
    authors = metadata.get("authors") or []
    keywords = metadata.get("keywords") or []
    keyword_text = ", ".join(keywords) if isinstance(keywords, list) else str(keywords)

    return {
        "Source File": metadata.get("source_file") or "N/A",
        "Title": metadata.get("title") or "N/A",
        "Authors": format_authors(authors) or "N/A",
        "Publication Date": metadata.get("publication_date") or "N/A",
        "Journal / Venue": metadata.get("journal_or_venue") or "N/A",
        "Volume": metadata.get("volume") or "N/A",
        "DOI": metadata.get("doi") or "N/A",
        "arXiv ID": metadata.get("arxiv_id") or "N/A",
        "Keywords": keyword_text or "N/A",
        "Language": metadata.get("language") or "N/A",
        "Quality Score": metadata.get("_quality_score", 0),
        "Abstract": metadata.get("abstract") or "N/A",
    }


def style_workbook(file_path: str) -> None:
    wb = load_workbook(file_path)
    ws = wb.active
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(COLUMNS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [18, 42, 42, 18, 42, 12, 30, 22, 42, 14, 14, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(file_path)


def append_metadata_to_excel(metadata: Dict[str, Any], excel_path: str = EXCEL_FILE_PATH) -> str:
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    row = flatten_metadata(metadata)
    df = pd.DataFrame([row], columns=COLUMNS)

    if not os.path.exists(excel_path):
        df.to_excel(excel_path, index=False)
    else:
        with pd.ExcelWriter(excel_path, mode="a", engine="openpyxl", if_sheet_exists="overlay") as writer:
            sheet_name = writer.book.active.title
            existing_rows = writer.book[sheet_name].max_row
            df.to_excel(writer, sheet_name=sheet_name, header=False, index=False, startrow=existing_rows)

    style_workbook(excel_path)
    return excel_path


def create_excel_from_metadata(metadata: Dict[str, Any], excel_path: str) -> str:
    """Create a fresh one-row workbook using the exact same 12-column schema."""
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    row = flatten_metadata(metadata)
    df = pd.DataFrame([row], columns=COLUMNS)
    df.to_excel(excel_path, index=False)
    style_workbook(excel_path)
    return excel_path
